#!/usr/bin/env python3
"""Audit the public PASTorIN treatment-level dataset and its source integrity."""

from __future__ import annotations

import hashlib
import json
import os
from collections import Counter, defaultdict
from pathlib import Path

import xlrd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "data/pastorin-treatment-comparisons"
MANIFEST = DATA / "manifest.json"
REPORT = ROOT / "docs/evidence/pastorin-treatment-comparisons-audit.json"


def digest(path: Path, algorithm: str) -> str:
    hasher = hashlib.new(algorithm)
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def is_present(value: object) -> bool:
    return value not in (None, "", "*")


def audit() -> dict:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    source_checks = []
    for entry in manifest["files"]:
        path = DATA / entry["name"]
        source_checks.append({
            "name": entry["name"],
            "bytes": path.stat().st_size,
            "archivedBytes": entry["archivedBytes"],
            "apiDeclaredBytes": entry["apiDeclaredBytes"],
            "apiDeclaredSizeMatchesDownload": entry["apiDeclaredBytes"] == path.stat().st_size,
            "md5": digest(path, "md5"),
            "expectedMd5": entry["md5"],
            "sha256": digest(path, "sha256"),
            "expectedSha256": entry["sha256"],
            "verified": (
                path.stat().st_size == entry["archivedBytes"]
                and digest(path, "md5") == entry["md5"]
                and digest(path, "sha256") == entry["sha256"]
            ),
        })
    if not all(item["verified"] for item in source_checks):
        raise ValueError("PASTorIN source file integrity check failed")

    dictionary_book = load_workbook(DATA / "PASTorIN_1_VARIABLES.xlsx", read_only=True, data_only=True)
    dictionary_rows = list(dictionary_book.active.iter_rows(min_row=2, values_only=True))
    dictionary_by_name = {row[1]: row for row in dictionary_rows if len(row) > 1 and row[1]}

    workbook = xlrd.open_workbook(DATA / "PASTorIN_2_DATA.xls", on_demand=True)
    if workbook.nsheets != 1:
        raise ValueError(f"Expected one data sheet, found {workbook.nsheets}")
    sheet = workbook.sheet_by_index(0)
    field_names = [str(sheet.cell_value(1, column)).strip() for column in range(39)]
    if len(field_names) != len(set(field_names)) or set(field_names) != set(dictionary_by_name):
        raise ValueError("Workbook field names do not match the archived dictionary")
    rows = []
    for row_index in range(2, sheet.nrows):
        rows.append({field: sheet.cell_value(row_index, column) for column, field in enumerate(field_names)})
    if len(rows) != 248:
        raise ValueError(f"Expected 248 treatment rows, found {len(rows)}")

    def counts(field: str, subset: list[dict] | None = None) -> dict:
        values = [row[field] for row in (subset or rows)]
        valid = [value for value in values if is_present(value)]
        return {"rows": len(values), "present": len(valid), "missing": len(values) - len(valid)}

    species = Counter(str(row["2_animal_species"]).strip().lower() for row in rows)
    environments = Counter(str(row["1-feeding_environment"]).strip().lower() for row in rows)
    cattle = [row for row in rows if str(row["2_animal_species"]).strip().lower() == "cattle"]
    cattle_stages = Counter(str(row["2_animal_physiological_stage"]).strip().lower() for row in cattle)
    cattle_environments = Counter(str(row["1-feeding_environment"]).strip().lower() for row in cattle)
    cattle_experiment_environments: dict[tuple[object, object], set[str]] = defaultdict(set)
    for row in cattle:
        cattle_experiment_environments[(row["1_publication_IDi"], row["1_experimentation_id"])].add(
            str(row["1-feeding_environment"]).strip().lower()
        )
    cattle_complete_adg_dmi = sum(
        is_present(row["2_animal_average_daily_gain"]) and is_present(row["4_diet_dry_matter_intake"])
        for row in cattle
    )

    report = {
        "classification": "public_literature_compilation_treatment_level_context_not_engine_validation",
        "source": {
            "persistentId": manifest["persistentId"],
            "version": manifest["apiVersion"],
            "releaseDate": manifest["releaseDate"],
            "repositoryLicense": manifest["repositoryLicense"],
            "dataPaper": {
                "doi": manifest["dataPaper"]["doi"],
                "articleLicense": manifest["dataPaper"]["articleLicense"],
                "datasetLicenseAsReportedInDataPaper": manifest["dataPaper"]["datasetLicenseAsReportedInDataPaper"],
            },
            "licenseNote": manifest["licenseNote"],
            "files": source_checks,
            "apiDeclaredSizeMismatch": any(not item["apiDeclaredSizeMatchesDownload"] for item in source_checks),
        },
        "workbook": {
            "sheetCount": workbook.nsheets,
            "sheetName": sheet.name,
            "rawRowsIncludingTwoHeaderRows": sheet.nrows,
            "rawColumns": sheet.ncols,
            "meaningfulVariableCount": len(field_names),
            "dictionaryMatchesWorkbook": True,
            "treatmentRows": len(rows),
            "unitOfRow": "treatment/group mean, not one animal",
            "individualAnimalIdentifierAvailable": False,
        },
        "coverage": {
            "species": dict(sorted(species.items())),
            "feedingEnvironment": dict(sorted(environments.items())),
            "cattleTreatmentRows": len(cattle),
            "cattleEnvironment": dict(sorted(cattle_environments.items())),
            "cattlePhysiologicalStage": dict(sorted(cattle_stages.items())),
            "cattleExperimentsWithBothPastureAndIndoorRows": sum(
                {"pasture", "indoor"}.issubset(values)
                for values in cattle_experiment_environments.values()
            ),
            "cattleExperimentsKeyedByPublicationAndExperiment": len(cattle_experiment_environments),
            "cattleRowsWithBothAdgAndDmi": cattle_complete_adg_dmi,
            "selectedFieldPresence": {
                field: counts(field) for field in (
                    "2_animals_number",
                    "2_animal_age",
                    "2_animal_initial_liveweight",
                    "2_animal_final_liveweight",
                    "2_animal_average_daily_gain",
                    "4_diet_dry_matter_intake",
                    "4_duration_of_daily_grazing",
                )
            },
            "cattleFieldPresence": {
                field: counts(field, cattle) for field in (
                    "2_animal_age",
                    "2_animal_initial_liveweight",
                    "2_animal_final_liveweight",
                    "2_animal_average_daily_gain",
                    "4_diet_dry_matter_intake",
                    "4_duration_of_daily_grazing",
                )
            },
        },
        "limitations": [
            "The 248 rows are group treatment means assembled from published papers; the workbook does not provide individual-animal observations or IDs for leakage-safe model evaluation.",
            "Only 11 cattle treatment rows contain both ADG and a DMI value; intake/digestibility measures are sparse and were reported under heterogeneous study methods and bases.",
            "The article describes a non-orthogonal, unbalanced multi-study compilation; pasture and indoor categories are not a globally randomized or paired counterfactual.",
            "The dataset has no aligned initial pasture stock, costs, or simulation state needed to evaluate the Agroville scenario end to end.",
            "The dataset paper reports 248 treatment rows from 93 annotated publications; this audit does not reconstruct or independently verify all source-paper extraction decisions.",
            "For the dictionary file, the repository API byte-size field disagrees with the original-format download even though the downloaded file matches the API MD5; both values are recorded for traceability.",
        ],
        "intendedUse": manifest["intendedUse"],
    }
    return report


if __name__ == "__main__":
    result = audit()
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    temporary_report = REPORT.with_name(f"{REPORT.name}.{os.getpid()}.tmp")
    temporary_report.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    os.replace(temporary_report, REPORT)
    print(json.dumps({key: result[key] for key in ("classification", "workbook", "coverage")}, indent=2))
