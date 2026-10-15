#!/usr/bin/env python3
"""Normalize the public Dryad mob/rotational pasture biomass tables.

The original XLSX and README remain immutable in data/raw. Row-level samples
and a predeclared paddock-level aggregation are exported separately. No model
is fitted and missing biomass values are not imputed.
"""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "data/raw/dryad-79cnp5ht4/version-6.zip"
OUT_DIR = ROOT / "data/dryad-mob-rotational-biomass/processed"
SAMPLE_PARQUET = OUT_DIR / "dryad-per-harvest-biomass-samples-v1.parquet"
PADDOCK_PARQUET = OUT_DIR / "dryad-paddock-biomass-height-v1.parquet"
REPORT = ROOT / "docs/evidence/dryad-mob-rotational-biomass-quality-2026-10.json"
EXPECTED_MEMBERS = {
    "Mob_&_Rotational_Study_Data_Summary.xlsx": "47c2aaf50572c8a46b8ab5e3dccd532d36c67fe04fd0f942c5c4f818d281b03b",
    "README_Mob_Rotational_Grazing.docx": "cf1215c422a74c17bfd48d0a26c46d6647964a44b50b253091edf618a7a0086e",
}


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def process() -> dict:
    archive_bytes = SOURCE.read_bytes()
    source_hash = sha256_bytes(archive_bytes)
    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        members = {Path(name).name: archive.read(name) for name in archive.namelist() if not name.endswith("/")}
    if set(members) != set(EXPECTED_MEMBERS):
        raise ValueError(f"unexpected Dryad archive members: {sorted(members)}")
    for name, expected_hash in EXPECTED_MEMBERS.items():
        actual_hash = sha256_bytes(members[name])
        if actual_hash != expected_hash:
            raise ValueError(f"Dryad internal file checksum mismatch for {name}: {actual_hash}")

    workbook = openpyxl.load_workbook(io.BytesIO(members["Mob_&_Rotational_Study_Data_Summary.xlsx"]), read_only=True, data_only=True)
    biomass_sheet = workbook["Per Harvest Biomass"]
    biomass_iter = biomass_sheet.iter_rows(values_only=True)
    biomass_header = next(biomass_iter)
    if biomass_header[:6] != ("Julian Day", "Plot", "Rep", "Year", "Graze", "PGFM Kg/ha"):
        raise ValueError(f"unexpected biomass columns: {biomass_header[:6]}")

    samples = []
    missing_target_rows = []
    group_values: dict[tuple[int, int, int, str], list[float]] = defaultdict(list)
    group_sample_n: Counter = Counter()
    for row_number, row in enumerate(biomass_iter, start=2):
        if not any(value is not None for value in row[:5]):
            continue
        julian_day, plot, replicate, year_code, treatment, raw_yield = row[:6]
        if not isinstance(julian_day, (int, float)) or not isinstance(plot, (int, float)):
            raise ValueError(f"invalid key at source row {row_number}: {row[:6]}")
        key = (int(year_code), int(julian_day), int(replicate), str(treatment))
        if not isinstance(raw_yield, (int, float)):
            missing_target_rows.append({
                "sourceRow": row_number,
                "julianDay": int(julian_day),
                "plotId": int(plot),
                "replicate": int(replicate),
                "studyYearCode": int(year_code),
                "treatment": str(treatment),
                "sourceValue": raw_yield,
            })
            continue
        value = float(raw_yield)
        if value < 0:
            raise ValueError(f"negative biomass at source row {row_number}")
        sample = {
            "sourceRow": row_number,
            "julianDay": int(julian_day),
            "plotId": int(plot),
            "replicate": int(replicate),
            "studyYearCode": int(year_code),
            "calendarYear": 2014 + int(year_code),
            "treatment": str(treatment),
            "preGrazingForageMassKgPerHa": value,
            "sourceTargetUnit": "kg/ha",
        }
        samples.append(sample)
        group_values[key].append(value)
        group_sample_n[key] += 1
    if len(missing_target_rows) != 17 or len(samples) != 623:
        raise ValueError(f"unexpected biomass completeness: {len(samples)} valid; {len(missing_target_rows)} missing")

    height_sheet = workbook["Pasture Heights"]
    height_iter = height_sheet.iter_rows(values_only=True)
    height_header = next(height_iter)
    if height_header[:6] != ("Pre/Post", "Julian Day", "Year", "Paddock/rep", "Graze", "Mean height (cm)"):
        raise ValueError(f"unexpected height columns: {height_header[:6]}")
    heights: dict[tuple[int, int, int, str], float] = {}
    for row_number, row in enumerate(height_iter, start=2):
        phase, julian_day, year_code, replicate, treatment, height = row[:6]
        if phase != "Pre":
            continue
        if not all(isinstance(value, (int, float)) for value in (julian_day, year_code, replicate, height)):
            raise ValueError(f"invalid pre-grazing height at source row {row_number}: {row[:6]}")
        key = (int(year_code), int(julian_day), int(replicate), str(treatment))
        if key in heights:
            raise ValueError(f"duplicate pre-grazing height key: {key}")
        heights[key] = float(height)
    if len(heights) != 64 or set(heights) != set(group_values):
        raise ValueError(f"height/biomass keys differ: {len(heights)} height keys vs {len(group_values)} biomass groups")

    paddocks = []
    for key, values in sorted(group_values.items()):
        year_code, julian_day, replicate, treatment = key
        paddocks.append({
            "paddockId": f"{treatment}-rep-{replicate}",
            "replicate": replicate,
            "treatment": treatment,
            "studyYearCode": year_code,
            "calendarYear": 2014 + year_code,
            "julianDay": julian_day,
            "nQuadratSamples": len(values),
            "preGrazingForageMassMeanKgPerHa": sum(values) / len(values),
            "preGrazingHeightMeanCm": heights[key],
            "aggregation": "arithmetic mean of available plot-level per-harvest biomass values within treatment x replicate x study year x Julian day",
        })
    if len(paddocks) != 64:
        raise ValueError(f"expected 64 paddock-time rows, got {len(paddocks)}")
    workbook.close()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    pq.write_table(pa.Table.from_pylist(samples), SAMPLE_PARQUET, compression="zstd")
    pq.write_table(pa.Table.from_pylist(paddocks), PADDOCK_PARQUET, compression="zstd")
    report = {
        "schemaVersion": "1.0.0",
        "status": "processed_candidate_external_biomass_cohort",
        "sourceDoi": "10.5061/dryad.79cnp5ht4",
        "sourceArchive": SOURCE.relative_to(ROOT).as_posix(),
        "sourceBytes": len(archive_bytes),
        "sourceSha256": source_hash,
        "officialInternalFileSha256": {name: sha256_bytes(members[name]) for name in sorted(members)},
        "rowLevelParquet": SAMPLE_PARQUET.relative_to(ROOT).as_posix(),
        "rowLevelSha256": sha256_bytes(SAMPLE_PARQUET.read_bytes()),
        "rowLevelRows": len(samples),
        "rowsWithMissingTargetNotExported": len(missing_target_rows),
        "missingTargetRows": missing_target_rows,
        "aggregatedParquet": PADDOCK_PARQUET.relative_to(ROOT).as_posix(),
        "aggregatedSha256": sha256_bytes(PADDOCK_PARQUET.read_bytes()),
        "aggregatedRows": len(paddocks),
        "groupCount": len({row["paddockId"] for row in paddocks}),
        "grouping": ["treatment", "replicate"],
        "repeatDimensions": ["calendarYear", "julianDay"],
        "quadratCountPerPaddockTime": {
            "minimum": min(row["nQuadratSamples"] for row in paddocks),
            "median": sorted(row["nQuadratSamples"] for row in paddocks)[len(paddocks) // 2],
            "maximum": max(row["nQuadratSamples"] for row in paddocks),
        },
        "units": {"outcome": "kg/ha (source label PGFM; interpreted as pre-grazing forage mass)", "predictor": "cm"},
        "featureOutcomeMatch": "Mean pre-grazing height at paddock/replicate level is matched to the arithmetic mean of available plot-level biomass quadrats within treatment, replicate, study year, and Julian day.",
        "limitations": [
            "17 of 640 source row slots contain missing PGFM and were not imputed.",
            "Paddock-level height does not match every individual quadrat; external evaluation must use the documented aggregate target.",
            "Only eight treatment x replicate paddocks recur over four years; repeated time points must remain grouped by paddock.",
            "This cohort supports a separate forage biomass evaluation only, not animal DMI, ADG, or the full engine.",
        ],
    }
    REPORT.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(process(), indent=2))
